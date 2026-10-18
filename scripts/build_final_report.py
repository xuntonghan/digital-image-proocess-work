#!/usr/bin/env python3
"""Build a polished Word report from the final project Markdown file."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BLUE = "17365D"
MID_BLUE = "366092"
PALE_BLUE = "EEF5FB"
GRAY = "666666"
LIGHT_GRAY = "D9E2F3"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LIGHT_GRAY, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, *, east_asia: str = "Noto Sans CJK SC", size: float | None = None) -> None:
    run.font.name = east_asia
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        run._element.rPr.rFonts.set(qn(f"w:{key}"), east_asia)
    if size is not None:
        run.font.size = Pt(size)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for level, size, color in ((1, 16, BLUE), (2, 13.5, BLUE), (3, 11.5, MID_BLUE), (4, 10.5, BLACK)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if level > 1 else 14)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Cm(0.75)
        style.paragraph_format.first_line_indent = Cm(-0.35)
        style.paragraph_format.space_after = Pt(3)

    code = doc.styles.add_style("Report Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Menlo"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    code.font.size = Pt(8)
    code.paragraph_format.left_indent = Cm(0.45)
    code.paragraph_format.right_indent = Cm(0.45)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(2)
    code.paragraph_format.line_spacing = 1.0
    code._element.get_or_add_pPr().append(OxmlElement("w:keepLines"))

    caption = doc.styles["Caption"]
    caption.font.name = "Noto Sans CJK SC"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(8)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.35)
        section.right_margin = Cm(2.1)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)
        section.different_first_page_header_footer = True


def add_header_footer(section, title: str) -> None:
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(title)
    set_run_font(run, east_asia="Noto Sans CJK SC", size=8.5)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document, title: str, subtitle: str, date: str) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, east_asia="Noto Sans CJK SC", size=25)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(subtitle)
    set_run_font(r, east_asia="Noto Sans CJK SC", size=14)
    r.font.color.rgb = RGBColor.from_string(MID_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("图像增强 · 几何纠正 · 图像压缩")
    set_run_font(r, east_asia="Noto Sans CJK SC", size=12)
    r.bold = True

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"实验日期：{date}")
    set_run_font(r, size=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("实验平台：NVIDIA GeForce RTX 4090 24GB / Apple MPS")
    set_run_font(r, size=10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("代码仓库：xuntonghan/digital-image-proocess-work（私有）")
    set_run_font(r, size=10)
    doc.add_page_break()


def parse_inline(paragraph, text: str) -> None:
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, east_asia="Noto Sans CJK SC", size=9)
            run.font.name = "Menlo"
            run.font.color.rgb = RGBColor.from_string(MID_BLUE)
        elif token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
        else:
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if match:
                run = paragraph.add_run(match.group(1))
                set_run_font(run)
                run.font.color.rgb = RGBColor.from_string(MID_BLUE)
                run.underline = True
            else:
                run = paragraph.add_run(token.replace("\\|", "|"))
                set_run_font(run)


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        values = [v.strip() for v in line.strip().strip("|").split("|")]
        rows.append(values)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", value) for value in rows[1]):
        rows.pop(1)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_idx, values in enumerate(rows):
        tr_pr = table.rows[row_idx]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for col_idx in range(width):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if row_idx == 0:
                set_cell_shading(cell, BLUE)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(0)
            value = values[col_idx] if col_idx < len(values) else ""
            parse_inline(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, size=8.2)
                if row_idx == 0:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
                    run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_image(doc: Document, root: Path, alt: str, path_text: str) -> None:
    path = (root / path_text).resolve() if not Path(path_text).is_absolute() else Path(path_text)
    if not path.exists():
        paragraph = doc.add_paragraph()
        parse_inline(paragraph, f"[缺失图片：{path_text}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(15.2))
    keep_with_next(paragraph)
    caption = doc.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    parse_inline(caption, alt)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph(style="Report Code")
    paragraph.paragraph_format.first_line_indent = Cm(0)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6F8")
    p_pr.append(shd)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Menlo"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
        run.font.size = Pt(7.5)
        if index < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def add_markdown(doc: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    root = markdown_path.parent
    index = 0
    code_lines: list[str] = []
    in_code = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line.rstrip())
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            add_image(doc, root, image.group(1), image.group(2))
            index += 1
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                index += 1
                continue
            paragraph = doc.add_paragraph(style=f"Heading {min(level - 1, 4)}")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            parse_inline(paragraph, text)
            keep_with_next(paragraph)
            index += 1
            continue
        if stripped == "---":
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet or numbered:
            paragraph = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
            paragraph.paragraph_format.first_line_indent = Cm(-0.35)
            parse_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line or next_line.startswith(("#", "|", "```", "![")):
                break
            if re.match(r"^([-*]\s+|\d+\.\s+)", next_line) or next_line == "---":
                break
            paragraph_lines.append(next_line)
            index += 1
        paragraph = doc.add_paragraph()
        parse_inline(paragraph, " ".join(paragraph_lines))


def build(source: Path, output: Path, date: str) -> None:
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    add_cover(
        doc,
        "数字图像处理课程实习",
        "最终实验报告与完整复现指南",
        date,
    )
    add_markdown(doc, source)
    for section in doc.sections:
        add_header_footer(section, "数字图像处理课程实习最终实验报告")
    core = doc.core_properties
    core.title = "数字图像处理课程实习最终实验报告"
    core.subject = "SIED 图像增强、GeoCorrection 几何纠正与 Kodak 图像压缩"
    core.author = "Han Xuntong"
    core.keywords = "数字图像处理, SIED, GeoCorrection, Kodak, RTX 4090"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--date", default="2026-10-18")
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve(), args.date)


if __name__ == "__main__":
    main()
